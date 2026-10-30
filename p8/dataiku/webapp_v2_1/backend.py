import io
import logging
import os
import re
import sys
import threading
import traceback
from pathlib import Path
from typing import Any, Dict, List, Optional

import pandas as pd
from flask import Response, jsonify, request

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    _XLSX_ERROR = None
except Exception as _exc_xlsx:
    Workbook = None
    _XLSX_ERROR = f"{type(_exc_xlsx).__name__}: {_exc_xlsx}"

try:
    _HERE = Path(__file__).resolve().parent
except NameError:
    _HERE = None

_PROJECT_ROOT = os.getenv("PII_PROJECT_ROOT")
_PATH_CANDIDATES = []
if _HERE is not None:
    _PATH_CANDIDATES.extend([_HERE.parent, _HERE])
if _PROJECT_ROOT:
    _PATH_CANDIDATES.append(Path(_PROJECT_ROOT))
for _candidate in _PATH_CANDIDATES:
    if _candidate.exists() and str(_candidate) not in sys.path:
        sys.path.insert(0, str(_candidate))

try:
    from pii_engine.adapters.dataiku_adapter import (
        DataikuAdapter,
        ModelArtifactError,
    )
    _ENGINE_IMPORT_ERROR = None
except Exception as _exc_engine:
    DataikuAdapter = None

    class ModelArtifactError(RuntimeError):
        pass

    _ENGINE_IMPORT_ERROR = f"{type(_exc_engine).__name__}: {_exc_engine}"

logger = logging.getLogger(__name__)

if "app" not in globals():
    from flask import Flask

    app = Flask(__name__)

_MAX_UPLOAD_MB = int(os.getenv("PII_MAX_UPLOAD_MB", "50"))
_MAX_READ_ROWS = int(os.getenv("PII_MAX_READ_ROWS", "50000"))
app.config["MAX_CONTENT_LENGTH"] = _MAX_UPLOAD_MB * 1024 * 1024

_adapter: Optional[Any] = None
_adapter_lock = threading.Lock()


def get_adapter() -> Any:
    global _adapter
    with _adapter_lock:
        if _adapter is None:
            if _ENGINE_IMPORT_ERROR:
                raise RuntimeError(
                    f"No se pudo importar pii_engine: {_ENGINE_IMPORT_ERROR}"
                )
            _adapter = DataikuAdapter(
                binary_folder=os.getenv("PII_BINARY_FOLDER", "model_binary_v1"),
                entity_folder=os.getenv("PII_ENTITY_FOLDER", "model_entity_v1"),
                content_folder=os.getenv("PII_CONTENT_FOLDER", "model_content_v1"),
                max_read_rows=_MAX_READ_ROWS,
                strict_models=True,
            )
        return _adapter


def set_adapter(adapter: Optional[Any]) -> None:
    global _adapter
    with _adapter_lock:
        _adapter = adapter


def _error(message: str, status: int = 400):
    return jsonify({"status": "error", "message": message}), status


def _parse_sample_size(value: Any) -> Optional[int]:
    if value is None or value == "":
        return None
    try:
        sample_size = int(value)
    except (TypeError, ValueError) as exc:
        raise ValueError("sample_size debe ser un entero") from exc
    if sample_size < 1 or sample_size > 50000:
        raise ValueError("sample_size debe estar entre 1 y 50000")
    return sample_size


def _read_uploaded_file(file, max_rows: int) -> pd.DataFrame:
    filename = file.filename or ""
    suffix = Path(filename).suffix.lower()

    if suffix == ".csv":
        try:
            return pd.read_csv(file, nrows=max_rows, low_memory=False)
        except pd.errors.EmptyDataError as exc:
            raise ValueError(f"'{filename}' está vacío.") from exc

    if suffix in {".xls", ".xlsx"}:
        try:
            return pd.read_excel(file, nrows=max_rows)
        except TypeError:
            df = pd.read_excel(file)
            return df.head(max_rows)
        except pd.errors.EmptyDataError as exc:
            raise ValueError(f"'{filename}' está vacío.") from exc

    raise ValueError(
        f"Formato no soportado para '{filename}'. Sube archivos .csv o .xlsx."
    )


def _prepare_dataframe(df: pd.DataFrame) -> pd.DataFrame:
    if df is None or df.shape[1] == 0:
        return df

    keep = []
    names = []
    used = set()
    for position, original in enumerate(df.columns):
        base = str(original).strip()
        if not base or base.lower() == "nan":
            continue
        name = base
        suffix = 2
        while name in used:
            name = f"{base}_{suffix}"
            suffix += 1
        used.add(name)
        keep.append(position)
        names.append(name)

    if not keep:
        return df.iloc[:, 0:0].copy()

    result = df.iloc[:, keep].copy()
    result.columns = names
    return result


def _result_to_row(result: Any, source_file: str) -> Dict[str, Any]:
    data = result.to_dict()
    is_pii = bool(data.get("pii"))

    return {
        "source_file": source_file,
        "dataset": source_file,
        "name": data.get("column"),
        "is_pii": is_pii,
        "pii_probability": data.get("prob_pii"),
        "entity": data.get("entity") if is_pii else None,
        "entity_probability": data.get("prob_entity") if is_pii else None,
        "needs_review": data.get("needs_review"),
        "evidence": data.get("evidence"),
        "match_rate": data.get("match_rate"),
        "sample_size": data.get("sample_size"),
        "model_version": data.get("model_version"),
    }


@app.route("/health", methods=["GET"])
def health():
    if _ENGINE_IMPORT_ERROR:
        return jsonify({
            "status": "degraded",
            "engine_ready": False,
            "error": _ENGINE_IMPORT_ERROR,
        }), 503

    try:
        info = get_adapter().health()
    except Exception as exc:
        logger.exception("Error al verificar el motor")
        return jsonify({
            "status": "degraded",
            "engine_ready": False,
            "error": f"{type(exc).__name__}: {exc}",
        }), 503

    status = 200 if info.get("engine_ready") else 503
    return jsonify(info), status


@app.route("/process_table", methods=["POST"])
def process_table():
    if _ENGINE_IMPORT_ERROR:
        return _error(
            f"No se pudo importar pii_engine: {_ENGINE_IMPORT_ERROR}",
            503,
        )

    files = request.files.getlist("file")
    if not files:
        return _error("No se subió ningún archivo.", 400)

    try:
        sample_size = _parse_sample_size(request.form.get("sample_size"))
    except ValueError as exc:
        return _error(str(exc), 400)

    try:
        adapter = get_adapter()
    except Exception as exc:
        logger.exception("No se pudo inicializar el adaptador")
        return _error(
            f"No se pudo inicializar el motor PII: {type(exc).__name__}: {exc}",
            503,
        )

    rows: List[Dict[str, Any]] = []
    warnings: List[str] = []

    for file in files:
        if not file.filename:
            continue

        try:
            df = _read_uploaded_file(file, _MAX_READ_ROWS)
        except ValueError as exc:
            warnings.append(str(exc))
            continue

        if df.shape[1] == 0:
            warnings.append(f"'{file.filename}' no tiene columnas para analizar.")
            continue

        df = _prepare_dataframe(df)
        if df.shape[1] == 0:
            warnings.append(f"'{file.filename}' no tiene columnas válidas para analizar.")
            continue

        try:
            results = adapter.analyze_dataframe(
                df,
                table_name=file.filename,
                sample_size=sample_size,
            )
        except ModelArtifactError as exc:
            logger.exception("Error de artefacto al analizar '%s'", file.filename)
            return _error(str(exc), 503)
        except Exception as exc:
            logger.exception("Error al analizar '%s'", file.filename)
            warnings.append(
                f"No se pudo analizar '{file.filename}': "
                f"{type(exc).__name__}: {exc}"
            )
            continue

        rows.extend(_result_to_row(result, file.filename) for result in results)

    if not rows:
        return _error(
            " | ".join(warnings) if warnings else "Ningún archivo pudo procesarse.",
            400,
        )

    response = {"status": "success", "columns_analysis": rows}
    if warnings:
        response["warnings"] = warnings
    return jsonify(response)


EXPORT_XLSX_HEADER = ["Archivo", "Nombre de Columna", "Prediccion",
                      "Probabilidad PII (%)", "Entidad", "Probabilidad Entidad (%)"]


def _safe_xlsx_filename(name):
    base = re.sub(r"[\\/:*?\"<>|]", "_", str(name or "reporte")).strip()[:80] or "reporte"
    return base if base.lower().endswith(".xlsx") else base + ".xlsx"


@app.route('/export_xlsx', methods=['POST'])
def export_xlsx():
    try:
        if Workbook is None:
            return jsonify({
                "status": "error",
                "message": f"openpyxl no disponible: {_XLSX_ERROR}. Usa CSV.",
            }), 500
        payload = request.get_json(force=True, silent=True) or {}
        rows = payload.get("rows") or []
        if not isinstance(rows, list) or not rows:
            return jsonify({"status": "error", "message": "Sin filas para exportar."}), 400

        wb = Workbook()
        ws = wb.active
        ws.title = "PII"
        ws.freeze_panes = "A2"
        head_fill = PatternFill("solid", fgColor="003746")
        head_font = Font(bold=True, color="FFFFFF")
        ws.append(EXPORT_XLSX_HEADER)
        for cell in ws[1]:
            cell.fill = head_fill
            cell.font = head_font
            cell.alignment = Alignment(vertical="center")
        for r in rows:
            vals = list(r)[:6] + [""] * max(0, 6 - len(r))
            out = []
            for idx, v in enumerate(vals[:6]):
                if idx in (3, 5) and v not in (None, ""):
                    try:
                        out.append(float(v))
                        continue
                    except (TypeError, ValueError):
                        pass
                out.append("" if v is None else str(v))
            ws.append(out)
        for col in ws.columns:
            width = max((len(str(c.value or "")) for c in col), default=10)
            ws.column_dimensions[col[0].column_letter].width = min(width + 2, 48)
        for row in ws.iter_rows(min_row=2, min_col=4, max_col=6):
            for cell in row:
                if isinstance(cell.value, float):
                    cell.number_format = '0.00'
        ws.auto_filter.ref = ws.dimensions

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return Response(
            buf.getvalue(),
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition":
                     f"attachment; filename=\"{_safe_xlsx_filename(payload.get('filename'))}\""},
        )
    except Exception as e:
        traceback.print_exc()
        return jsonify({"status": "error", "message": str(e)}), 500
